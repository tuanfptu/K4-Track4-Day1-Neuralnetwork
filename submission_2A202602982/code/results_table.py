"""results_table.py — Lưu kết quả ra JSON và tự động điền vào experiments.xlsx từ template.

Tên cột của sheet "Experiments":
    exp_id, group, description, loss, optimizer, lr, weight_decay, batch, epochs, hidden, dropout,
    clip_norm, precision, init, seed, step0_loss, best_val_loss, best_epoch, final_train_loss,
    final_val_loss, val_acc, val_macro_f1, time_per_epoch_s, peak_mem_MB, diverged,
    eval_acc, eval_macro_f1, figure_file, notes
"""
from __future__ import annotations

import json
from pathlib import Path
import openpyxl


def save_result(result: dict, results_dir: str = "../results") -> str:
    """Ghi result['cfg'], result['history'], result['summary'] (không ghi best_state) ra JSON."""
    out_dir = Path(results_dir)
    out_dir.mkdir(parents=True, exist_ok=True)
    exp_id = result["cfg"]["exp_id"]
    file_path = out_dir / f"{exp_id}.json"

    data_to_save = {
        "cfg": result["cfg"],
        "history": result["history"],
        "summary": result["summary"],
    }

    with open(file_path, "w", encoding="utf-8") as f:
        json.dump(data_to_save, f, indent=2, ensure_ascii=False)

    return str(file_path)


def load_results(results_dir: str = "../results") -> list[dict]:
    """Đọc mọi file *.json trong results_dir, sắp theo exp_id."""
    r_dir = Path(results_dir)
    if not r_dir.exists():
        return []
    results = []
    for p in sorted(r_dir.glob("*.json")):
        with open(p, "r", encoding="utf-8") as f:
            results.append(json.load(f))
    return results


def to_row(result: dict, eval_scores: dict | None = None, notes: str = "") -> dict:
    """Biến một kết quả thành một dict dòng cho bảng experiments.xlsx."""
    cfg = result["cfg"]
    summary = result.get("summary", {})
    exp_id = cfg["exp_id"]

    hidden_val = cfg.get("hidden", (256, 128))
    if isinstance(hidden_val, (list, tuple)):
        hidden_str = "-".join(map(str, hidden_val))
    else:
        hidden_str = str(hidden_val)

    clip_val = cfg.get("clip_norm")
    clip_str = "none" if clip_val is None else clip_val

    eval_acc = None
    eval_macro_f1 = None
    if eval_scores is not None:
        eval_acc = eval_scores.get("eval_acc", eval_scores.get("acc"))
        eval_macro_f1 = eval_scores.get("eval_macro_f1", eval_scores.get("macro_f1"))

    return {
        "exp_id": exp_id,
        "group": cfg.get("group", "other"),
        "description": cfg.get("description", ""),
        "loss": cfg.get("loss", "ce").upper(),
        "optimizer": cfg.get("optimizer", ""),
        "lr": cfg.get("lr"),
        "weight_decay": cfg.get("weight_decay", 0.0),
        "batch": cfg.get("batch", 512),
        "epochs": cfg.get("epochs", 20),
        "hidden": hidden_str,
        "dropout": cfg.get("dropout", 0.0),
        "clip_norm": clip_str,
        "precision": cfg.get("precision", "fp32"),
        "init": cfg.get("init", "he"),
        "seed": cfg.get("seed", 1),
        "step0_loss": summary.get("step0_loss"),
        "best_val_loss": summary.get("best_val_loss"),
        "best_epoch": summary.get("best_epoch"),
        "final_train_loss": summary.get("final_train_loss"),
        "final_val_loss": summary.get("final_val_loss"),
        "val_acc": summary.get("val_acc"),
        "val_macro_f1": summary.get("val_macro_f1"),
        "time_per_epoch_s": summary.get("time_per_epoch_s"),
        "peak_mem_MB": summary.get("peak_mem_MB"),
        "diverged": summary.get("diverged", False),
        "eval_acc": eval_acc,
        "eval_macro_f1": eval_macro_f1,
        "figure_file": f"figures/{exp_id}.png",
        "notes": notes or cfg.get("notes", ""),
    }


def write_xlsx(rows: list[dict], template_path: str, out_path: str,
               summary_notes: dict[str, str] | None = None) -> None:
    """Điền các dòng vào sheet Experiments của mẫu, cập nhật công thức và lưu ra out_path."""
    wb = openpyxl.load_workbook(template_path)
    ws = wb["Experiments"]

    header_cols = {}
    for col_idx in range(1, ws.max_column + 1):
        h = ws.cell(1, col_idx).value
        if h:
            header_cols[h] = col_idx

    start_row = 2
    for i, r_data in enumerate(rows):
        row_num = start_row + i

        for key, val in r_data.items():
            if key in header_cols:
                col_idx = header_cols[key]
                ws.cell(row=row_num, column=col_idx, value=val)

        # Thiết lập các cột công thức cho hàng này
        r = row_num
        ws.cell(row=r, column=header_cols["step0_gap_vs_lnC"], value=f'=IF(P{r}="","",P{r}-LN(7))')
        ws.cell(row=r, column=header_cols["gap_val_minus_train"], value=f'=IF(OR(T{r}="",S{r}=""),"",T{r}-S{r})')
        ws.cell(row=r, column=header_cols["delta_val_f1_vs_base"], value=f'=IF(OR(V{r}="",Seeds!$C$8=""),"",V{r}-Seeds!$C$8)')
        ws.cell(row=r, column=header_cols["beyond_noise"], value=f'=IF(OR(AF{r}="",Seeds!$C$10=""),"",IF(ABS(AF{r})>Seeds!$C$10,"Có","Không"))')

    # Điền nhận xét vào sheet Summary nếu có
    if summary_notes and "Summary" in wb.sheetnames:
        ws_sum = wb["Summary"]
        for r_idx in range(2, 12):
            grp_name = ws_sum.cell(r_idx, 1).value
            if grp_name in summary_notes:
                ws_sum.cell(r_idx, 8, value=summary_notes[grp_name])

    out_file = Path(out_path)
    out_file.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out_file)
    print(f"Đã lưu bảng thí nghiệm ({len(rows)} dòng) -> {out_file}")
